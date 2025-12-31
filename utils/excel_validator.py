# utils/excel_validator.py

from openpyxl import load_workbook


def is_excel_valid(expected_hash: str) -> bool:
    try:
        wb = load_workbook("testcases/manual/testcases.xlsx")
        ws = wb["TestCases"]

        stored_hash = ws["B1"].value
        if stored_hash != expected_hash:
            return False

        # ✅ Ensure at least ONE test case exists
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[0] is not None:
                return True

        return False  # No test cases found

    except Exception:
        return False
