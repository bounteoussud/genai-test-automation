from openpyxl import load_workbook


def has_test_cases(excel_path: str) -> bool:
    try:
        wb = load_workbook(excel_path)
        ws = wb["TestCases"]

        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[0] is not None:
                return True

        return False
    except Exception:
        return False
